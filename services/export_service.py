"""
Export Service.

This service provides functionality to export data in various formats
(CSV, Excel) for brand analytics and reports.
"""

import csv
import io
from typing import List, Dict, Any
from datetime import datetime
import logging

logger = logging.getLogger(__name__)

try:
    import openpyxl
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    EXCEL_AVAILABLE = True
except ImportError:
    EXCEL_AVAILABLE = False
    logger.warning("openpyxl not available. Excel export will be disabled.")


class ExportService:
    """Service for exporting data to various formats."""
    
    def export_to_csv(self, data: List[Dict], filename: str = None) -> str:
        """
        Export data to CSV format.
        
        Args:
            data (list): List of dictionaries to export
            filename (str, optional): Output filename
            
        Returns:
            str: CSV content as string
        """
        if not data:
            return ""
        
        output = io.StringIO()
        fieldnames = data[0].keys()
        
        writer = csv.DictWriter(output, fieldnames=fieldnames)
        writer.writeheader()
        writer.writerows(data)
        
        csv_content = output.getvalue()
        output.close()
        
        if filename:
            try:
                with open(filename, 'w', encoding='utf-8', newline='') as f:
                    f.write(csv_content)
                logger.info(f"Exported CSV to {filename}")
            except Exception as e:
                logger.error(f"Error writing CSV file: {e}")
        
        return csv_content
    
    def export_to_excel(self, data: List[Dict], filename: str = None, sheet_name: str = "Data") -> bytes:
        """
        Export data to Excel format.
        
        Args:
            data (list): List of dictionaries to export
            filename (str, optional): Output filename
            sheet_name (str): Sheet name
            
        Returns:
            bytes: Excel file content
        """
        if not EXCEL_AVAILABLE:
            raise ImportError("openpyxl is required for Excel export. Install it with: pip install openpyxl")
        
        if not data:
            wb = Workbook()
            ws = wb.active
            ws.title = sheet_name
            ws['A1'] = "No data available"
        else:
            wb = Workbook()
            ws = wb.active
            ws.title = sheet_name
            
            # Write headers
            headers = list(data[0].keys())
            for col_idx, header in enumerate(headers, 1):
                cell = ws.cell(row=1, column=col_idx, value=header)
                cell.font = Font(bold=True)
                cell.fill = PatternFill(start_color="CCCCCC", end_color="CCCCCC", fill_type="solid")
                cell.alignment = Alignment(horizontal="center")
            
            # Write data
            for row_idx, row_data in enumerate(data, 2):
                for col_idx, header in enumerate(headers, 1):
                    value = row_data.get(header, "")
                    ws.cell(row=row_idx, column=col_idx, value=value)
            
            # Auto-adjust column widths
            for col in ws.columns:
                max_length = 0
                col_letter = col[0].column_letter
                for cell in col:
                    try:
                        if len(str(cell.value)) > max_length:
                            max_length = len(str(cell.value))
                    except:
                        pass
                adjusted_width = min(max_length + 2, 50)
                ws.column_dimensions[col_letter].width = adjusted_width
        
        # Save to file or return bytes
        if filename:
            wb.save(filename)
            logger.info(f"Exported Excel to {filename}")
            with open(filename, 'rb') as f:
                return f.read()
        else:
            output = io.BytesIO()
            wb.save(output)
            output.seek(0)
            return output.read()
    
    def export_sentiment_data(self, brand_name: str, format: str = 'csv') -> Any:
        """
        Export sentiment data for a brand.
        
        Args:
            brand_name (str): Brand name
            format (str): 'csv' or 'excel'
            
        Returns:
            str or bytes: Exported data
        """
        try:
            from services.brand_analytics_service import BrandAnalyticsService
            
            brand_service = BrandAnalyticsService()
            
            # Get product performance data
            product_performance = brand_service.get_product_performance(brand_name)
            
            # Format for export
            export_data = []
            for product in product_performance:
                export_data.append({
                    'Product Name': product['product_name'],
                    'Category': product['category'],
                    'Total Reviews': product['total_reviews'],
                    'Positive Count': product['positive_count'],
                    'Negative Count': product['negative_count'],
                    'Neutral Count': product['neutral_count'],
                    'Positive %': product['positive_pct'],
                    'Average Rating': product['average_rating']
                })
            
            if format.lower() == 'excel':
                return self.export_to_excel(export_data, sheet_name=f"{brand_name} Sentiment")
            else:
                return self.export_to_csv(export_data)
                
        except Exception as e:
            logger.error(f"Error exporting sentiment data for {brand_name}: {e}")
            raise
    
    def export_trend_report(self, brand_name: str, format: str = 'excel') -> Any:
        """
        Export trend report for a brand.
        
        Args:
            brand_name (str): Brand name
            format (str): 'csv' or 'excel'
            
        Returns:
            str or bytes: Exported report
        """
        try:
            from services.trend_service import TrendService
            
            trend_service = TrendService()
            report = trend_service.generate_trend_report(brand_name, time_period='month')
            
            if format.lower() == 'excel':
                wb = Workbook()
                
                # Sheet 1: Summary
                ws1 = wb.active
                ws1.title = "Summary"
                ws1['A1'] = f"Trend Report: {brand_name}"
                ws1['A2'] = f"Generated: {datetime.utcnow().strftime('%Y-%m-%d %H:%M:%S')}"
                ws1['A4'] = "Sentiment Metrics"
                ws1['A5'] = "Total Reviews"
                ws1['B5'] = report['sentiment_metrics']['total_reviews']
                ws1['A6'] = "Positive %"
                ws1['B6'] = f"{report['sentiment_metrics']['positive_pct']}%"
                ws1['A7'] = "Negative %"
                ws1['B7'] = f"{report['sentiment_metrics']['negative_pct']}%"
                ws1['A8'] = "Average Rating"
                ws1['B8'] = report['sentiment_metrics']['average_rating']
                
                # Sheet 2: Product Performance
                ws2 = wb.create_sheet("Product Performance")
                if report['product_performance']:
                    headers = list(report['product_performance'][0].keys())
                    for col_idx, header in enumerate(headers, 1):
                        ws2.cell(row=1, column=col_idx, value=header)
                    for row_idx, product in enumerate(report['product_performance'], 2):
                        for col_idx, header in enumerate(headers, 1):
                            ws2.cell(row=row_idx, column=col_idx, value=product.get(header, ""))
                
                # Sheet 3: Common Concerns
                ws3 = wb.create_sheet("Common Concerns")
                ws3['A1'] = "Concern"
                ws3['B1'] = "Frequency"
                for row_idx, concern in enumerate(report['common_concerns'], 2):
                    ws3.cell(row=row_idx, column=1, value=concern['concern'])
                    ws3.cell(row=row_idx, column=2, value=concern['frequency'])
                
                output = io.BytesIO()
                wb.save(output)
                output.seek(0)
                return output.read()
            else:
                # CSV format - combine all data
                export_data = []
                for product in report['product_performance']:
                    export_data.append({
                        'Type': 'Product',
                        'Name': product['product_name'],
                        'Metric': 'Performance',
                        'Value': product['positive_pct']
                    })
                for concern in report['common_concerns']:
                    export_data.append({
                        'Type': 'Concern',
                        'Name': concern['concern'],
                        'Metric': 'Frequency',
                        'Value': concern['frequency']
                    })
                return self.export_to_csv(export_data)
                
        except Exception as e:
            logger.error(f"Error exporting trend report for {brand_name}: {e}")
            raise

